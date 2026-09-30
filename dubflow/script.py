"""English lines + voice evidence -> Mongolian dub script with character names and eleven_v3 tags.

Each line reaches the model with its pyannote voice id, the gender/age/emotion heard on it, and what that
voice was called in earlier chunks, so speaker names stay tied to real voices across the whole film.
Chunks are cached in cache/llm/ keyed by their full prompt, so re-runs are free.
"""
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

from pydantic import BaseModel

from .analyze import age_group, gender_word
from .common import FFMPEG, ROOT, TAG

MODEL = "gpt-6-luna"
CHUNK = 40  # lines per request
MIN_CHUNK = 5  # a failing chunk is halved down to this size before giving up
CONTEXT = 15  # previous translated lines shown with each chunk
LOOKAHEAD = 8  # following lines shown, so sentences split over lines are translated whole
CHARS_PER_SEC = 15  # Mongolian characters that fit one second of dub at normal pace
EMOTION_MIN = 0.5  # emotion hints below this confidence are left out
AUDIO_PAD = 1.0  # Gemini: seconds of audio kept around each chunk
CACHE_DIR = ROOT / "cache" / "llm"

TAGS = ("quietly, whispers, softly, sad, crying, sobbing, hurt, bitter, sarcastic, angry, furious, shouting, "
        "screaming, excited, happy, playful, laughs, chuckles, sighs, gasps, nervous, worried, scared, trembling, "
        "confused, curious, surprised, shocked, hesitant, determined, serious, urgent, calm, gentle, reflective, "
        "cold, teasing, pleading, firm, warm, relieved, hopeful, tired, embarrassed")
# Speaker's own sounds that may also stay as tags.
SOUNDS = "giggles, laughing, sniffs, sniffles, grunts, groans, coughs, clears throat, panting, gulps, screams, sobs"
ALLOWED_TAGS = {t.strip() for t in (TAGS + ", " + SOUNDS).split(",")}
# Tags the model invents anyway -> nearest allowed one, so the line keeps its delivery instead of going flat.
TAG_MAP = {"concerned": "worried", "anxious": "worried", "tense": "nervous", "fearful": "scared",
           "panicked": "scared", "reassuring": "gentle", "sweet": "warm", "tender": "warm", "grateful": "warm",
           "sincere": "warm", "earnest": "serious", "grave": "serious", "somber": "sad", "emotional": "sad",
           "weary": "tired", "exhausted": "tired", "cheerful": "happy", "amused": "chuckles", "mocking": "sarcastic",
           "dry": "sarcastic", "threatening": "cold", "commanding": "firm", "formal": "calm", "polite": "calm",
           "professional": "calm", "matter-of-fact": "calm", "neutral": "calm", "apologetic": "hesitant",
           "frustrated": "angry", "annoyed": "angry", "irritated": "angry", "desperate": "pleading",
           "pained": "hurt", "breathless": "panting"}
FRAGMENT_SEC = 1.5  # max length of a piece join_fragments may hand back to the surrounding speaker

SYSTEM = """You write the Mongolian dub script for an English-dubbed {genre} ("{title}").
{premise}

For every input line return: id (unchanged), speaker, text.

Evidence on each line:
- voice: the diarized voice id from the real audio. The same voice id is almost always the same character for
  the whole film; a voice actor may double a few minor roles. "voice_history" tells what each voice was called
  in earlier lines: keep that name unless the context clearly shows a different character.
- heard: gender and age group measured from the voice ("uncertain" = no clear call), plus the emotion the audio
  sounds like when it was clear. Gender is reliable: never give a line heard as female to a male character or
  the other way round. The emotion model is trained on acted speech and over-reports "sad": trust the words
  and context first.

speaker:
- The character who says the line, in English. Reuse known character names. Narration is usually {narrator}.
- For unnamed minor roles use a short English role that includes the gender ("Village Woman 1",
  "Soldier Man", "Merchant Man") and keep it the same across lines.

text: the Mongolian dub line.
{style}
- Keep names English-style, spelled exactly: {names}. Other names: transliterate the English form and reuse
  the spelling from earlier lines.
- Fit the time slot: stay near the line's max_chars (shorter is fine, much longer is not).
- Sentences split over several lines: "next_lines" shows what follows. Translate the whole sentence, then
  cut the Mongolian at a natural point so each line still carries its part (Mongolian puts the verb last, so
  moving words between the pieces is fine). Every piece must still sound like speech, not a broken fragment.
  If a piece clearly continues the previous speaker's sentence, give it that speaker.
- Fix obvious transcription errors from context.

Tags (for the TTS voice):
- Start almost every line with one or two performance tags in English square brackets that describe how THIS
  speaker says it, e.g. [warm], [sad] [quietly], [angry] [shouting], [firm], [excited]. Allowed tags: {tags}.
  A tag can go mid-line when the emotion changes. Only a truly flat line goes without one.
- The speaker's own sounds in the input ([laughs], [sighs], [gasps], [sobbing]) become tags in the right place.
- Background sounds are NOT the speaker: drop [crowd murmuring], [birds chirping], [music], [door opens] etc.
- A line with only the speaker's own sound stays tag-only, e.g. "[laughs]"."""

STYLE = """- Sound like real Mongolians talking to each other today, the way a good Mongolian TV dub sounds - not a
  translation, not a book. Rephrase freely; keep the meaning, intent, emotion and who-is-addressed exactly, but
  never the English word order, idioms or sentence build.
- Everyday spoken Khalkha: short sentences, drop pronouns when obvious, and the particles and spoken endings a
  Mongolian really uses: шүү, шүү дээ, дээ, л, л дээ, биз дээ, юм, юм чинь, байхгүй юу, байх аа, шиг байна,
  тэ, аа/ээ/оо; -чихлаа, -чихсан, -чих, -ъя, -гаарай, -даа, -хгүй юу, -сан юм уу.
- Live-speech openers and reactions where the feeling calls for them: за, заа, өө, аан, хөөе, ёо, яахав,
  нээрээ юу, тэгэлгүй яахав, үгүй ээ, болъё доо, юу гэнээ?, яасан бэ?, одоо яах юм. Emotion can show as
  repetition or a broken start ("Би... би мэдээгүй"). Don't stuff them in every line - one where a person would.
- Not literary: in dialogue no -жээ/-чээ/-лээ шүү past ("мартжээ" -> "мартчихсан", "хэлжээ" -> "хэлсэн"), no -вч
  ("харагдавч" -> "харагдаж байгаа ч"), no мэт/төдий/нэгэн цагт/дэндүү/шахуу/авч одох -> шиг/л/урьд нь/хэтэрхий/
  бараг/авах. No bookish/official words (хэрэгжүүлэх, эрхэм, тухайлбал) unless the character is formal.
  Narration ("I/we" looking back) may be a little softer and storytelling, but still said aloud, not written.
- Examples: "You really are cruel." -> "Ямар хатуу хүн бэ чи." | "Did you eat lunch?" -> "Хоолоо идсэн үү?" |
  "I had told Noah too many gentle lies." -> "Би Ноад хэтэрхий олон зөөлөн худал хэлсэн." | "Are you serious?" ->
  "Нээрээ юу?" | "Don't worry." -> "Санаа зовох хэрэггүй ээ." | "Fine, let's go." -> "За за, явъя."
- Use Mongolian words, not Russian/English loans, when a common Mongolian word exists: отряд -> бүлэг/цэрэг,
  лагерь -> хуаран, бизнес -> наймаа/ажил, проблем -> асуудал, окей -> за, идея -> санаа.
- Grammar must be correct: case endings, vowel harmony, possessives (-аа/-ээ/-оо, чинь, минь).
- Politeness: та to elders, parents-in-law, superiors and strangers; чи between spouses, siblings, friends
  and to children. Family address words as a Mongolian would say them (ээж, аав, ах, эгч, дүү, хүү минь).
- Numbers and abbreviations as Mongolian words (2025 -> хоёр мянга хорин тав), no digits.
- Period/fantasy setting: no modern slang or modern-world words."""

POLISH = """You are the native Mongolian script editor of a TV dub ("{title}", {genre}). Each line has the English
original, who says it, the gender/age heard on the voice, and a draft Mongolian dub line. Return every line
(same id) with the final speaker and text.

Speaker: keep it unless it is clearly wrong: nobody thanks, calls or answers themselves by name ("Thank you,
Isola" is not said by Isola); a short piece that just continues someone's unfinished sentence belongs to that
person; a line heard as female is not said by a male character and the other way round. Only use speaker
names that already appear in the lines or previous_lines.

Fix the draft so a Mongolian viewer hears natural, correct, living speech:
- Meaning: must match the English (intent, facts, who is addressed). Fix mistranslations and lost meaning.
- Language: correct grammar, case endings and vowel harmony; replace calques, bookish words and needless
  Russian/English loans with what a Mongolian would really say.
{style}
- Sentences split over consecutive lines must read as one natural sentence when played in order.
- Keep each line near its max_chars (the dub must fit the time slot); shorten by dropping filler, not meaning.
- Read each line aloud in your head: if it sounds like a subtitle or a book, rewrite it the way that character
  would really say it in that moment.
- Keep names exactly as in the draft. Keep the square-bracket performance tags (they drive the TTS emotion); fix
  one only if it is wrong for the line; a background sound is not a tag. Tags stay in English and must be from:
  {tags}.
If a draft line is already good, return it unchanged."""

AUDIO_NOTE = """
The audio of these lines is attached; each line's "at" is its start in seconds from the clip start.
Listen to it to confirm who speaks and how each line is delivered."""


class Line(BaseModel):
    id: int
    speaker: str
    text: str


class Chunk(BaseModel):
    lines: list[Line]


class Fixed(BaseModel):
    id: int
    speaker: str
    text: str


class FixedChunk(BaseModel):
    lines: list[Fixed]


def read_guide(path):
    """Plain text of a character/story guide (.docx or .txt)."""
    if not path or not Path(path).exists():
        return ""
    path = Path(path)
    if path.suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        rows = [" | ".join(c.text.strip() for c in r.cells) for t in d.tables for r in t.rows]
        return "\n".join([p.text for p in d.paragraphs if p.text.strip()] + rows)
    return path.read_text(encoding="utf-8")


def heard(line):
    if "p_female" not in line:
        return None
    parts = [gender_word(line["p_female"]), age_group(line["age"])]
    if line.get("emotion_score", 0) >= EMOTION_MIN:
        parts.append(line["emotion"])
    return ", ".join(parts)


def audio_clip(audio, start, end):
    return subprocess.run([FFMPEG, "-v", "quiet", "-ss", f"{start:.2f}", "-t", f"{end - start:.2f}", "-i", str(audio),
                           "-ac", "1", "-ar", "16000", "-b:a", "32k", "-f", "mp3", "-"],
                          capture_output=True, check=True).stdout


def _client(model):
    if model.startswith("gemini"):
        from google import genai
        return genai.Client()
    from openai import OpenAI
    return OpenAI()


def _ask(client, model, system, user, clip, schema):
    if model.startswith("gemini"):
        from google.genai import types
        parts = ([types.Part.from_bytes(data=clip, mime_type="audio/mp3")] if clip else []) + [user]
        r = client.models.generate_content(
            model=model, contents=parts,
            config=types.GenerateContentConfig(system_instruction=system, response_mime_type="application/json",
                                               response_schema=schema))
        return r.parsed, r.usage_metadata.total_token_count
    r = client.responses.parse(model=model, reasoning={"effort": "medium"}, text_format=schema,
                               input=[{"role": "system", "content": system}, {"role": "user", "content": user}])
    return r.output_parsed, r.usage.total_tokens


def _cached(client, model, system, user, schema, n, clip=None):
    """Model call cached by its full prompt; retried until it returns exactly ids 0..n-1.
    Returns (None, tokens) if it keeps failing, so the caller can split the chunk."""
    path = CACHE_DIR / f"{hashlib.sha1('|'.join([model, system, user]).encode()).hexdigest()[:16]}.json"
    if path.exists():
        return schema.model_validate_json(path.read_text(encoding="utf-8")).lines, 0
    total = 0
    for attempt in range(2):
        parsed, used = _ask(client, model, system, user, clip, schema)
        total += used
        result = parsed.lines if parsed else []
        if sorted(l.id for l in result) == list(range(n)):
            path.write_text(schema(lines=result).model_dump_json(indent=1), encoding="utf-8")
            return result, total
        print(f"      got {len(result)}/{n} lines back, retrying ({attempt + 1}/2)", flush=True)
    return None, total


def _chunks(total, size):
    """(start, length) of each chunk, as a queue: a chunk the model keeps failing is re-queued as two halves."""
    return [(n, min(size, total - n)) for n in range(0, total, size)]


def _split(queue, n, size):
    if size <= MIN_CHUNK:
        sys.exit(f"model keeps failing on lines {n}-{n + size}; check them in lines.json")
    half = size // 2
    print(f"      splitting lines {n}-{n + size} in two", flush=True)
    queue[:0] = [(n, half), (n + half, size - half)]


def _film_notes(film):
    notes = f"\n\n{film['notes']}" if film.get("notes") else ""
    if film.get("glossary"):
        notes += f"\n\nGlossary (use these Mongolian terms): {json.dumps(film['glossary'], ensure_ascii=False)}"
    if film.get("address"):
        notes += f"\n\nHow characters address each other: {film['address']}"
    return notes


def max_chars(line):
    return max(10, int((line["end"] - line["start"]) * CHARS_PER_SEC))


def write_script(lines, film, voices, model, audio=None, chunk_size=CHUNK, polish_model=None):
    """Returns lines with "speaker", "draft" (first pass) and "mn" (after the polish pass) added."""
    guide = read_guide(ROOT / film["guide"] if film.get("guide") else None)
    system = SYSTEM.format(title=film.get("title", "untitled"), genre=film.get("genre", "drama"),
                           premise=film.get("premise", ""), narrator=film.get("narrator", "the main character"),
                           names=json.dumps(film.get("names", {}), ensure_ascii=False), tags=TAGS, style=STYLE)
    system += _film_notes(film)
    system += f"\n\nCaption spellings of characters: {json.dumps(film['aliases'], ensure_ascii=False)}" \
        if film.get("aliases") else ""
    system += f"\n\nCharacter and story guide:\n{guide}" if guide else ""
    client = _client(model)
    if model.startswith("gemini"):
        system += AUDIO_NOTE

    not_speakers = set(film.get("not_speakers", []))
    characters = sorted((set(film.get("names", {})) | set(film.get("aliases", {}))) - not_speakers)
    history = defaultdict(Counter)  # voice -> Counter(character)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    done, tokens = [], 0
    queue = _chunks(len(lines), chunk_size)
    while queue:
        n, size = queue.pop(0)
        chunk = lines[n:n + size]
        clip_start = max(0.0, chunk[0]["start"] - AUDIO_PAD)
        in_chunk = sorted({l["voice"] for l in chunk} - {"SFX"})
        payload = {
            "known_characters": characters,
            "voice_history": {v: {"voice": f"{voices[v].get('gender', '?')}, {voices[v].get('age_group', '?')}",
                                  "called_so_far": dict(history[v].most_common(4))}
                              for v in in_chunk if v in voices},
            "previous_lines": [{"speaker": d["speaker"], "voice": d["voice"], "en": d["en"], "mn": d["draft"]}
                               for d in done[-CONTEXT:]],
            "lines": [{"id": i, "voice": l["voice"], "heard": heard(l), "seconds": round(l["end"] - l["start"], 1),
                       "max_chars": max_chars(l), "en": l["en"]}
                      | ({"at": round(l["start"] - clip_start, 1)} if audio else {})
                      for i, l in enumerate(chunk)],
            "next_lines": [l["en"] for l in lines[n + size:n + size + LOOKAHEAD]],
        }
        user = json.dumps(payload, ensure_ascii=False, indent=1)
        clip = audio_clip(audio, clip_start, chunk[-1]["end"] + AUDIO_PAD) if audio else None
        result, used = _cached(client, model, system, user, Chunk, len(chunk), clip)
        tokens += used
        if result is None:
            _split(queue, n, size)
            continue
        for r in sorted(result, key=lambda r: r.id):
            line = {**chunk[r.id], "speaker": r.speaker.strip(), "draft": r.text.strip()}
            done.append(line)
            if line["voice"] != "SFX":
                history[line["voice"]][line["speaker"]] += 1
        characters = sorted(set(characters) | {d["speaker"] for d in done})
        print(f"      draft  [{n + size}/{len(lines)}] "
              f"{'cached' if not used else f'{used} tokens'}", flush=True)

    tokens += polish(done, film, polish_model or model, chunk_size)
    join_fragments(done)
    print(f"      {tokens} tokens this run")
    return done


def _tag(m):
    tag = m.group(0)[1:-1].strip().lower()
    tag = TAG_MAP.get(tag, tag)
    return f"[{tag}]" if tag in ALLOWED_TAGS else ""


def clean_tags(text):
    """Keep only tags the TTS should perform (near-synonyms mapped onto allowed ones); background sounds
    ([crowd murmuring], [birds chirping]) go."""
    return re.sub(r"\s{2,}", " ", TAG.sub(_tag, text)).strip()


def join_fragments(lines):
    """A short piece sandwiched inside one speaker's unfinished sentence is that speaker's (diarization jitter)."""
    for a, b, c in zip(lines, lines[1:], lines[2:]):
        if (a["speaker"] == c["speaker"] != b["speaker"] and b["end"] - b["start"] < FRAGMENT_SEC
                and a["en"].rstrip()[-1:] not in ".?!" and b["start"] - a["end"] < 0.5):
            b["speaker"] = a["speaker"]


def polish(lines, film, model, chunk_size=CHUNK):
    """Second pass: a native-editor rewrite of the draft for meaning, grammar and natural speech. Sets "mn"."""
    system = POLISH.format(title=film.get("title", "untitled"), genre=film.get("genre", "drama"), style=STYLE,
                           tags=TAGS + ", " + SOUNDS)
    system += _film_notes(film)
    client, tokens = _client(model), 0
    queue = _chunks(len(lines), chunk_size)
    while queue:
        n, size = queue.pop(0)
        chunk = lines[n:n + size]
        payload = {
            "previous_lines": [{"speaker": d["speaker"], "mn": d["mn"]} for d in lines[max(0, n - CONTEXT):n]],
            "lines": [{"id": i, "speaker": l["speaker"], "heard": heard(l), "max_chars": max_chars(l), "en": l["en"],
                       "draft": l["draft"]} for i, l in enumerate(chunk)],
            "next_lines": [{"speaker": l["speaker"], "en": l["en"], "draft": l["draft"]}
                           for l in lines[n + size:n + size + LOOKAHEAD]],
        }
        result, used = _cached(client, model, system, json.dumps(payload, ensure_ascii=False, indent=1),
                               FixedChunk, len(chunk))
        tokens += used
        if result is None:
            _split(queue, n, size)
            continue
        for r in result:
            chunk[r.id]["mn"] = clean_tags(r.text)
            chunk[r.id]["speaker"] = r.speaker.strip() or chunk[r.id]["speaker"]
        print(f"      polish [{n + size}/{len(lines)}] "
              f"{'cached' if not used else f'{used} tokens'}", flush=True)
    return tokens


def gender_warnings(lines, cast):
    """Lines whose heard gender contradicts their character's overall gender (likely misattributed)."""
    out = []
    for l in lines:
        c = cast.get(l["speaker"], {})
        g = gender_word(l["p_female"]) if "p_female" in l else "uncertain"
        if g != "uncertain" and c.get("gender") not in (None, "uncertain", g):
            out.append(f"{l['id']:>5}  {l['start']:8.1f}s  {l['speaker']} ({c['gender']}) but heard {g}: {l['en'][:70]}")
    return out
